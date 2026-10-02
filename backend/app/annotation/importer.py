"""Import sentence-level annotations made in the team's spreadsheet into the AnnotationStore.

Expected sheet columns: doc_id, case_name, sentence_id, sentence_text, label, annotator, notes
(labels FACT / LAW / PRECEDENT / ARGUMENT / RULING). Each document becomes one case:

- its sentences, in sentence_id order, are pinned as the case's segmentation (the team labelled
  these exact sentences, so our segmenter must not re-split them);
- their concatenation is the judgment text, screened and stored by JudgmentCollector like any
  other judgment;
- each row becomes a LabeledSentence (source "manual").

Cases already present in the store are skipped, so re-running an import is harmless.
"""

import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.annotation.collector import JudgmentCollector
from app.annotation.label_scheme import LabelScheme
from app.annotation.models import LabeledSentence, sentence_id
from app.annotation.store import AnnotationStore

LABEL_MAP: dict[str, str] = {
    "FACT": "Facts",
    "FACTS": "Facts",
    "LAW": "Law Applied",
    "LAW APPLIED": "Law Applied",
    "PRECEDENT": "Precedent",
    "ARGUMENT": "Argument",
    "RULING": "Ruling",
}
COLUMNS = ("doc_id", "case_name", "sentence_id", "sentence_text", "label", "annotator", "notes")
_YEAR_IN_TITLE = re.compile(r"\b(19[5-9]\d|20\d\d)\b")


@dataclass
class Row:
    doc_id: str
    case_name: str
    position: float
    text: str
    label: str
    annotator: str
    notes: str


@dataclass
class CaseResult:
    case_id: str
    title: str
    sentences: int
    status: str  # "imported" | "skipped" | "rejected"
    detail: str = ""
    label_counts: dict[str, int] = field(default_factory=dict)


@dataclass
class ImportReport:
    cases: list[CaseResult]
    problems: list[str]  # row-level issues that were skipped, e.g. unknown labels

    @property
    def imported(self) -> list[CaseResult]:
        return [c for c in self.cases if c.status == "imported"]


def normalize_doc_id(doc_id: Any) -> str:
    """'case 11', ' case11 ', 'Case11' -> 'case11'."""
    return re.sub(r"[^a-z0-9]+", "", str(doc_id).lower()) or "doc"


def read_rows(path: str | Path, sheet: str = "Annotations") -> tuple[list[Row], list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"sheet {sheet!r} not found; sheets are {wb.sheetnames}")
    raw = list(wb[sheet].iter_rows(values_only=True))
    header = [str(c).strip().lower() if c is not None else "" for c in raw[0]]
    missing = [c for c in COLUMNS[:5] if c not in header]
    if missing:
        raise ValueError(f"missing columns {missing}; header is {header}")
    col = {name: header.index(name) for name in COLUMNS if name in header}

    def cell(r: tuple, name: str) -> Any:
        i = col.get(name)
        return r[i] if i is not None and i < len(r) else None

    rows: list[Row] = []
    problems: list[str] = []
    for n, r in enumerate(raw[1:], start=2):
        if not any(c not in (None, "") for c in r):
            continue
        if str(cell(r, "sentence_id")).strip().lower() == "sentence_id":
            problems.append(f"row {n}: repeated header row skipped")
            continue
        text, label = str(cell(r, "sentence_text") or "").strip(), str(cell(r, "label") or "").strip().upper()
        try:
            position = float(cell(r, "sentence_id"))
        except (TypeError, ValueError):
            problems.append(f"row {n}: sentence_id {cell(r, 'sentence_id')!r} is not a number; row skipped")
            continue
        if not text:
            problems.append(f"row {n}: empty sentence_text; row skipped")
            continue
        if label not in LABEL_MAP:
            problems.append(f"row {n}: unknown label {label!r}; row skipped")
            continue
        rows.append(Row(
            doc_id=normalize_doc_id(cell(r, "doc_id")),
            case_name=str(cell(r, "case_name") or "").strip(),
            position=position,
            text=text,
            label=LABEL_MAP[label],
            annotator=str(cell(r, "annotator") or "").strip(),
            notes=str(cell(r, "notes") or "").strip(),
        ))
    return rows, problems


def import_annotations(
    path: str | Path,
    collector: JudgmentCollector,
    store: AnnotationStore,
    default_annotator: str,
    dry_run: bool = False,
) -> ImportReport:
    rows, problems = read_rows(path)
    by_doc: dict[str, list[Row]] = defaultdict(list)
    for row in rows:
        by_doc[row.doc_id].append(row)

    existing = set(store.case_ids())
    results: list[CaseResult] = []
    for case_id, doc_rows in by_doc.items():
        doc_rows.sort(key=lambda r: r.position)
        # most common case_name: a stray value in one row (e.g. "state") must not become the title
        names = Counter(r.case_name for r in doc_rows if r.case_name)
        title = names.most_common(1)[0][0] if names else ""
        positions = [r.position for r in doc_rows]
        if len(set(positions)) != len(positions):
            results.append(CaseResult(case_id, title, len(doc_rows), "rejected", "duplicate sentence_id within the case"))
            continue
        if case_id in existing:
            results.append(CaseResult(case_id, title, len(doc_rows), "skipped", "already in the annotation store"))
            continue

        sentences = [r.text for r in doc_rows]
        record = collector.ingest_text("\n".join(sentences), case_id=case_id, source_path=f"{Path(path).name}#{case_id}")
        record.title = title
        if m := _YEAR_IN_TITLE.findall(title):
            record.decision_year = int(m[-1])
        ok, reason = collector.is_eligible(record)
        if not ok:
            results.append(CaseResult(case_id, title, len(doc_rows), "rejected", reason))
            continue

        counts: dict[str, int] = defaultdict(int)
        for r in doc_rows:
            counts[r.label] += 1
        gaps = int(positions[-1] - positions[0] + 1) - len(positions)
        detail = f"{gaps} sentence_id gap(s) in the sheet" if gaps else ""
        results.append(CaseResult(case_id, title, len(doc_rows), "imported", detail, dict(counts)))
        if dry_run:
            continue
        collector.collect(record)
        store.save_segments(case_id, sentences, segmenter=f"xlsx:{Path(path).name}")
        store.append_many([
            LabeledSentence(
                sentence_id=sentence_id(case_id, idx),
                case_id=case_id,
                idx=idx,
                text=r.text,
                label=r.label,
                source="manual",
                annotator=r.annotator or default_annotator,
            )
            for idx, r in enumerate(doc_rows)
        ])
    assert all(LabelScheme.validate(v) for v in LABEL_MAP.values())
    return ImportReport(results, problems)
