"""A review sheet for the benchmark's 90 held-out test items (benchmark/labelling/review_test_90.xlsx).

The 60 scenarios and 30 out-of-scope questions in the seed-7 sample decide every scenario and
out-of-scope number in the paper, and their labels were written by an AI. This sheet lets a legal
reviewer check them:
- the question;
- the acceptable IPC sections, with titles;
- what LegalGPT (gpt-5.4-mini, final system) cited, so disagreements stand out;
- columns to fill in: agree yes/no, corrected sections, note.

Usage (from the repo root):
    python backend/scripts/make_review_sheet.py
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from app.benchmark.metrics import load_jsonl  # noqa: E402
from app.config import REPO_ROOT, get_settings  # noqa: E402
from app.core.corpus import load_chunks  # noqa: E402
from run_baselines import OUT, sample  # noqa: E402

OUT_FILE = REPO_ROOT / "benchmark" / "labelling" / "review_test_90.xlsx"
REFERENCE_RUN = "legalgpt_openai_gpt-5.4-mini_final"
HEAD = PatternFill("solid", fgColor="DDE7F5")
FILL = PatternFill("solid", fgColor="FFF8DC")  # cells the reviewer fills in


def titles() -> dict[str, str]:
    out: dict[str, str] = {}
    for c in load_chunks(get_settings().corpus_chunks_path):
        if c.section and c.title and c.citation_path.endswith(f"Section {c.section}"):
            out.setdefault(c.section, c.title.strip().rstrip("."))
    return out


def sheet(wb, name: str, header: list[str], rows: list[list[str]], widths: list[int], fill_from: int) -> None:
    ws = wb.create_sheet(name)
    ws.append(header)
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True), HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.column >= fill_from:
                cell.fill = FILL
    yes_no = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    ws.add_data_validation(yes_no)
    col = chr(64 + fill_from)  # the "agree?" column
    yes_no.add(f"{col}2:{col}{len(rows) + 1}")
    ws.freeze_panes = "B2"


def main() -> int:
    t = titles()
    items = sample(7, 40, 20, 60, 30)
    ref_path = OUT / f"{REFERENCE_RUN}.jsonl"
    cited = {r["id"]: sorted(set(r["extra_cited"])) for r in load_jsonl(ref_path)} if ref_path.exists() else {}
    named = lambda secs: "\n".join(f"{s} - {t.get(s, '?')}" for s in secs)

    wb = Workbook()
    intro = wb.active
    intro.title = "How to review"
    for line in [
        "LegalGPT benchmark: review of the 90 held-out test items",
        "",
        "These 90 items decide every scenario and out-of-scope number in the paper. Their labels were written by an AI.",
        "",
        "Sheet 'Scenarios' (60): is the list of acceptable IPC sections right for this situation?",
        "   - agree = yes if every listed section is a reasonable answer and no obviously right section is missing.",
        "   - agree = no: write the corrected list in 'corrected sections' (numbers separated by spaces, e.g. 415 420).",
        "   - 'LegalGPT cited' is for reference only: it shows what the system answered (gpt-5.4-mini, final run).",
        "",
        "Sheet 'Out of scope' (30): is this question really NOT about an offence under the IPC?",
        "   - agree = yes if it is a civil / consumer / family / other-law matter; no if it does describe an IPC offence.",
        "",
        "Fill in the reviewer name, role (e.g. law student, advocate) and date once on each sheet's first row.",
        "Save the file and give it back; the results are imported and every score is recomputed.",
    ]:
        intro.append([line])
    intro.column_dimensions["A"].width = 120
    intro["A1"].font = Font(bold=True, size=13)

    scen = [it for it in items if it["set"] == "scenario"]
    oos = [it for it in items if it["set"] == "oos"]
    sheet(wb, "Scenarios",
          ["id", "question", "acceptable IPC sections (AI-written label)", "LegalGPT cited", "agree? (yes/no)",
           "corrected sections", "note", "reviewer name, role, date"],
          [[it["id"], it["question"], named(it["acceptable"]), named(cited.get(it["id"], [])), "", "", "", ""] for it in scen],
          [9, 55, 45, 40, 12, 18, 30, 22], fill_from=5)
    sheet(wb, "Out of scope",
          ["id", "question", "LegalGPT cited (should be nothing)", "agree it is NOT an IPC offence? (yes/no)", "note",
           "reviewer name, role, date"],
          [[it["id"], it["question"], named(cited.get(it["id"], [])), "", "", ""] for it in oos],
          [9, 70, 30, 18, 35, 22], fill_from=4)
    OUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_FILE)
    print(f"wrote {OUT_FILE} ({len(scen)} scenarios, {len(oos)} out-of-scope)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
